from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.db.models import Detection

MONTHS_ES = (
    "ene",
    "feb",
    "mar",
    "abr",
    "may",
    "jun",
    "jul",
    "ago",
    "sep",
    "oct",
    "nov",
    "dic",
)


def _parse_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _as_aware(at: datetime) -> datetime:
    if at.tzinfo is None:
        return at.replace(tzinfo=timezone.utc)
    return at


def _format_absolute_date(at: datetime) -> str:
    local = _as_aware(at).astimezone()
    return f"{local.day} {MONTHS_ES[local.month - 1]} {local.year}, {local.strftime('%H:%M')}"


def _detection_label(event: Detection) -> str:
    return "Sana" if event.clase == "Sana" else "Antracnosis"


def _detection_status(event: Detection) -> str:
    if event.clase == "Sana":
        return event.estado or "saludable"
    return event.estado or event.severity or "leve"


def _operator_name(event: Detection) -> str:
    user = event.user
    if user is None:
        return "—"
    return user.full_name or user.username or "—"


def _serialize_report_row(event: Detection) -> dict:
    return {
        "id": event.id,
        "code": f"#DET-{event.id:04d}",
        "label": _detection_label(event),
        "origin": event.source,
        "status": _detection_status(event),
        "confidence": round(event.confidence * 100),
        "date": _format_absolute_date(event.created_at),
        "recomendacion": event.recomendacion,
        "operator": _operator_name(event),
    }


def query_report_detections(
    db: Session,
    *,
    date_from: str | None = None,
    date_to: str | None = None,
    clase: str | None = None,
    status: str | None = None,
) -> list[Detection]:
    stmt = (
        select(Detection)
        .options(joinedload(Detection.user))
        .where(Detection.is_active.is_(True))
        .order_by(Detection.created_at.desc())
    )

    start = _parse_date(date_from)
    end = _parse_date(date_to)
    if start is not None:
        start_dt = datetime.combine(start, time.min, tzinfo=timezone.utc)
        stmt = stmt.where(Detection.created_at >= start_dt)
    if end is not None:
        end_dt = datetime.combine(end, time.max, tzinfo=timezone.utc)
        stmt = stmt.where(Detection.created_at <= end_dt)

    if clase in ("Sana", "Antracnosis"):
        stmt = stmt.where(Detection.clase == clase)

    events = list(db.scalars(stmt).unique().all())

    if status in ("saludable", "leve", "moderado", "crítico"):
        events = [event for event in events if _detection_status(event) == status]

    return events


def build_report_summary(events: list[Detection]) -> dict:
    total = len(events)
    healthy = sum(1 for e in events if e.clase == "Sana")
    anthracnose = sum(1 for e in events if e.clase == "Antracnosis")
    mild = sum(1 for e in events if _detection_status(e) == "leve")
    moderate = sum(1 for e in events if _detection_status(e) == "moderado")
    severe = sum(1 for e in events if _detection_status(e) == "crítico")
    avg_confidence = (
        sum(e.confidence for e in events) / total * 100 if total > 0 else 0.0
    )

    return {
        "total": total,
        "healthy": healthy,
        "anthracnose": anthracnose,
        "severity": {
            "leve": mild,
            "moderada": moderate,
            "severa": severe,
        },
        "avg_confidence": round(avg_confidence, 1),
    }


def get_detections_report(
    db: Session,
    *,
    date_from: str | None = None,
    date_to: str | None = None,
    clase: str | None = None,
    status: str | None = None,
    limit: int = 500,
) -> dict:
    events = query_report_detections(
        db,
        date_from=date_from,
        date_to=date_to,
        clase=clase,
        status=status,
    )
    capped = max(1, min(limit, 2000))
    preview = events[:capped]
    return {
        "summary": build_report_summary(events),
        "detections": [_serialize_report_row(event) for event in preview],
        "returned": len(preview),
        "total_matching": len(events),
    }


def _style_header(ws, columns: int) -> None:
    fill = PatternFill("solid", fgColor="0F291E")
    font = Font(color="FFFFFF", bold=True)
    for col in range(1, columns + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws) -> None:
    for idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 48)


def build_detections_workbook(events: list[Detection]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Detecciones"

    headers = [
        "Código",
        "Fecha",
        "Clase",
        "Estado",
        "Confianza (%)",
        "Origen",
        "Operador",
        "Recomendación",
    ]
    ws.append(headers)
    _style_header(ws, len(headers))

    for event in events:
        ws.append(
            [
                f"#DET-{event.id:04d}",
                _format_absolute_date(event.created_at),
                _detection_label(event),
                _detection_status(event),
                round(event.confidence * 100),
                event.source,
                _operator_name(event),
                event.recomendacion or "",
            ]
        )

    _autosize(ws)

    summary = build_report_summary(events)
    ws_sum = wb.create_sheet("Resumen")
    ws_sum.append(["Métrica", "Valor"])
    _style_header(ws_sum, 2)
    rows = [
        ("Total detecciones", summary["total"]),
        ("Naranjas sanas", summary["healthy"]),
        ("Con antracnosis", summary["anthracnose"]),
        ("Severidad leve", summary["severity"]["leve"]),
        ("Severidad moderada", summary["severity"]["moderada"]),
        ("Severidad severa", summary["severity"]["severa"]),
        ("Confianza promedio (%)", summary["avg_confidence"]),
        ("Generado el", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M")),
    ]
    for label, value in rows:
        ws_sum.append([label, value])
    _autosize(ws_sum)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_detections_excel(
    db: Session,
    *,
    date_from: str | None = None,
    date_to: str | None = None,
    clase: str | None = None,
    status: str | None = None,
) -> tuple[bytes, str]:
    events = query_report_detections(
        db,
        date_from=date_from,
        date_to=date_to,
        clase=clase,
        status=status,
    )
    content = build_detections_workbook(events)
    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M")
    filename = f"reporte_detecciones_{stamp}.xlsx"
    return content, filename
